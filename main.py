"""
TxtFolderToExcel - Converts a folder of .txt files into a single Excel workbook.
- One sheet per .txt file (Line Number + Content table)
- Smart sorts sheets by detected dates (including "February 2nd 2026" style) or numbers
- First sheet is a clickable Table of Contents
- Back-to-TOC link on every sheet
"""

import os
import re
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from pathlib import Path
from datetime import datetime

# --- Check for openpyxl ---
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:
    import sys
    root = tk.Tk()
    root.withdraw()
    messagebox.showerror(
        "Missing Dependency",
        "This program requires the 'openpyxl' library.\n\n"
        "Install it by opening Command Prompt and running:\n\n"
        "    pip install openpyxl"
    )
    sys.exit(1)


# ---------- Smart Sorting ----------
MONTHS = {
    "january": 1, "jan": 1,
    "february": 2, "feb": 2,
    "march": 3, "mar": 3,
    "april": 4, "apr": 4,
    "may": 5,
    "june": 6, "jun": 6,
    "july": 7, "jul": 7,
    "august": 8, "aug": 8,
    "september": 9, "sept": 9, "sep": 9,
    "october": 10, "oct": 10,
    "november": 11, "nov": 11,
    "december": 12, "dec": 12,
}

# Numeric date patterns (most specific first)
NUMERIC_DATE_PATTERNS = [
    ("%Y-%m-%d", r"\b(\d{4}-\d{2}-\d{2})\b"),
    ("%Y_%m_%d", r"\b(\d{4}_\d{2}_\d{2})\b"),
    ("%Y%m%d",   r"\b(\d{8})\b"),
    ("%m-%d-%Y", r"\b(\d{2}-\d{2}-\d{4})\b"),
    ("%m_%d_%Y", r"\b(\d{2}_\d{2}_\d{4})\b"),
    ("%m-%d-%y", r"\b(\d{2}-\d{2}-\d{2})\b"),
    ("%m_%d_%y", r"\b(\d{2}_\d{2}_\d{2})\b"),
]

# Matches: "February 2nd 2026", "Feb 2 2026", "February 2, 2026",
#          "2nd February 2026", "2 Feb 2026", etc.
MONTH_NAMES_RE = (
    r"(?:January|February|March|April|May|June|July|August|September|"
    r"October|November|December|Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sept|Sep|Oct|Nov|Dec)"
)

# "Month Day Year"  e.g. February 2nd 2026 / Feb 2, 2026
MONTH_DAY_YEAR_RE = re.compile(
    rf"\b({MONTH_NAMES_RE})\s+(\d{{1,2}})(?:st|nd|rd|th)?[,\s]+(\d{{4}})\b",
    re.IGNORECASE,
)

# "Day Month Year"  e.g. 2nd February 2026 / 2 Feb 2026
DAY_MONTH_YEAR_RE = re.compile(
    rf"\b(\d{{1,2}})(?:st|nd|rd|th)?\s+({MONTH_NAMES_RE})[,\s]+(\d{{4}})\b",
    re.IGNORECASE,
)


def try_parse_date(name: str):
    """Return a datetime if a date can be found in the filename, else None."""
    # 1) Month name formats first (these are the most user-friendly)
    m = MONTH_DAY_YEAR_RE.search(name)
    if m:
        month_str, day_str, year_str = m.group(1), m.group(2), m.group(3)
        month = MONTHS.get(month_str.lower())
        try:
            if month:
                return datetime(int(year_str), month, int(day_str))
        except ValueError:
            pass

    m = DAY_MONTH_YEAR_RE.search(name)
    if m:
        day_str, month_str, year_str = m.group(1), m.group(2), m.group(3)
        month = MONTHS.get(month_str.lower())
        try:
            if month:
                return datetime(int(year_str), month, int(day_str))
        except ValueError:
            pass

    # 2) Numeric date formats
    for fmt, pattern in NUMERIC_DATE_PATTERNS:
        for match in re.findall(pattern, name):
            try:
                return datetime.strptime(match, fmt)
            except ValueError:
                continue

    return None


def natural_sort_key(name: str):
    """
    Returns a tuple that sorts files intelligently:
      1) Files with dates first (sorted chronologically)
      2) Then files with numbers (sorted numerically: file2 < file10)
      3) Then everything else alphabetically (case-insensitive)
    """
    dt = try_parse_date(name)
    if dt is not None:
        return (0, dt.timestamp(), name.lower())

    parts = re.split(r"(\d+)", name)
    if any(p.isdigit() for p in parts):
        key = tuple((int(p) if p.isdigit() else p.lower()) for p in parts if p != "")
        return (1, key, name.lower())

    return (2, name.lower())


# ---------- Excel Helpers ----------
def sanitize_sheet_name(name: str) -> str:
    """Excel sheet names: max 31 chars, cannot contain \\ / ? * [ ] :"""
    for ch in ['\\', '/', '?', '*', '[', ']', ':']:
        name = name.replace(ch, '_')
    return name[:31].strip() or "Sheet"


# ---------- Main Generation ----------
def generate_excel(input_folder: str, output_file: str, status_label: tk.Label, progress: ttk.Progressbar):
    txt_files = list(Path(input_folder).glob("*.txt"))

    if not txt_files:
        messagebox.showwarning("No Files Found", "No .txt files were found in the selected folder.")
        return

    # Smart sort
    txt_files.sort(key=lambda p: natural_sort_key(p.stem))

    wb = Workbook()
    wb.remove(wb.active)

    # Styling
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    header_align = Alignment(horizontal="center", vertical="center")
    thin = Side(border_style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    link_font = Font(color="0563C1", underline="single")
    back_link_font = Font(color="0563C1", underline="single", italic=True)

    # ---- Create Table of Contents first ----
    toc = wb.create_sheet(title="Table of Contents")

    progress["maximum"] = len(txt_files) + 1
    used_names = set()
    sheet_info = []

    # ---- Pre-pass: build sheet names ----
    for txt_path in txt_files:
        base = sanitize_sheet_name(txt_path.stem)
        name = base
        suffix = 1
        while name.lower() in used_names:
            suffix += 1
            name = f"{base[:28]}_{suffix}"
        used_names.add(name.lower())
        sheet_info.append({"sheet_name": name, "path": txt_path})

    # ---- Build each data sheet ----
    for idx, info in enumerate(sheet_info, start=1):
        txt_path = info["path"]
        sheet_name = info["sheet_name"]

        status_label.config(text=f"Processing ({idx}/{len(sheet_info)}): {txt_path.name}")
        status_label.update_idletasks()

        ws = wb.create_sheet(title=sheet_name)

        # Back-to-TOC link at the very top
        back_cell = ws.cell(row=1, column=1, value="← Back to Table of Contents")
        back_cell.hyperlink = "#'Table of Contents'!A1"
        back_cell.font = back_link_font
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)

        # Headers (row 2)
        ws.cell(row=2, column=1, value="Line Number")
        ws.cell(row=2, column=2, value="Content")
        for col in (1, 2):
            cell = ws.cell(row=2, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border

        # Read file
        line_count = 0
        try:
            with open(txt_path, "r", encoding="utf-8", errors="replace") as f:
                lines = f.read().splitlines()
            line_count = len(lines)
        except Exception as e:
            ws.cell(row=3, column=1, value="ERROR")
            ws.cell(row=3, column=2, value=f"Could not read file: {e}")
            lines = []

        for i, line in enumerate(lines, start=1):
            ws.cell(row=i + 2, column=1, value=i).border = border
            c = ws.cell(row=i + 2, column=2, value=line)
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=False)

        # Excel Table (rows 2 → last)
        last_row = max(line_count + 2, 3)
        table_ref = f"A2:B{last_row}"
        safe_table_name = re.sub(r'\W+', '_', f"tbl_{idx}_{sheet_name}")[:50]
        tbl = Table(displayName=safe_table_name, ref=table_ref)
        tbl.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(tbl)

        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 100
        ws.freeze_panes = "A3"

        # Save info for TOC
        dt = try_parse_date(txt_path.stem)
        sheet_info[idx - 1]["line_count"] = line_count
        sheet_info[idx - 1]["detected_date"] = dt.strftime("%Y-%m-%d") if dt else ""

        progress["value"] = idx
        progress.update_idletasks()

    # ---- Fill in Table of Contents ----
    status_label.config(text="Building Table of Contents...")
    status_label.update_idletasks()

    toc["A1"] = "Table of Contents"
    toc["A1"].font = Font(bold=True, size=16, color="305496")
    toc.merge_cells("A1:E1")
    toc["A1"].alignment = Alignment(horizontal="left", vertical="center")

    toc["A2"] = f"Source folder: {input_folder}"
    toc["A2"].font = Font(italic=True, color="595959")
    toc.merge_cells("A2:E2")

    headers = ["#", "Sheet Name", "Original Filename", "Detected Date", "Line Count"]
    for col_idx, h in enumerate(headers, start=1):
        c = toc.cell(row=4, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = border

    for i, info in enumerate(sheet_info, start=1):
        row = 4 + i
        toc.cell(row=row, column=1, value=i).border = border

        link_cell = toc.cell(row=row, column=2, value=info["sheet_name"])
        link_cell.hyperlink = f"#'{info['sheet_name']}'!A1"
        link_cell.font = link_font
        link_cell.border = border

        toc.cell(row=row, column=3, value=info["path"].name).border = border
        toc.cell(row=row, column=4, value=info.get("detected_date", "")).border = border
        toc.cell(row=row, column=5, value=info.get("line_count", 0)).border = border

    # Format TOC as an Excel Table
    toc_last_row = 4 + len(sheet_info)
    toc_table = Table(displayName="TOC_Table", ref=f"A4:E{toc_last_row}")
    toc_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    toc.add_table(toc_table)

    # Column widths
    toc.column_dimensions["A"].width = 6
    toc.column_dimensions["B"].width = 34
    toc.column_dimensions["C"].width = 50
    toc.column_dimensions["D"].width = 16
    toc.column_dimensions["E"].width = 12
    toc.freeze_panes = "A5"

    # Make TOC the active sheet on open
    wb.active = wb.index(toc)

    progress["value"] = progress["maximum"]
    progress.update_idletasks()

    # ---- Save ----
    try:
        wb.save(output_file)
    except PermissionError:
        messagebox.showerror(
            "Save Failed",
            f"Could not save:\n{output_file}\n\n"
            "The file may be open in Excel. Please close it and try again."
        )
        return
    except Exception as e:
        messagebox.showerror("Save Failed", f"Error saving file:\n{e}")
        return

    status_label.config(text="Done!")
    messagebox.showinfo(
        "Success",
        f"Excel file created successfully!\n\n"
        f"Sheets created: {len(sheet_info)} (+ Table of Contents)\n"
        f"Output: {output_file}"
    )


# ---------- GUI ----------
class TxtToExcelApp:
    def __init__(self, root):
        self.root = root
        root.title("TxtFolderToExcel")
        root.geometry("640x270")
        root.resizable(False, False)

        pad = {"padx": 10, "pady": 6}

        tk.Label(root, text="Input Folder (containing .txt files):", anchor="w").grid(
            row=0, column=0, columnspan=3, sticky="w", **pad)
        self.input_var = tk.StringVar()
        tk.Entry(root, textvariable=self.input_var, width=68).grid(row=1, column=0, columnspan=2, **pad)
        tk.Button(root, text="Browse...", width=12, command=self.browse_input).grid(row=1, column=2, **pad)

        tk.Label(root, text="Output Excel File (Save As):", anchor="w").grid(
            row=2, column=0, columnspan=3, sticky="w", **pad)
        self.output_var = tk.StringVar()
        tk.Entry(root, textvariable=self.output_var, width=68).grid(row=3, column=0, columnspan=2, **pad)
        tk.Button(root, text="Save As...", width=12, command=self.browse_output).grid(row=3, column=2, **pad)

        self.generate_btn = tk.Button(
            root, text="Generate", width=20, height=2,
            bg="#305496", fg="white", font=("Segoe UI", 10, "bold"),
            command=self.on_generate
        )
        self.generate_btn.grid(row=4, column=0, columnspan=3, pady=12)

        self.progress = ttk.Progressbar(root, orient="horizontal", length=580, mode="determinate")
        self.progress.grid(row=5, column=0, columnspan=3, padx=10, pady=(0, 4))
        self.status_label = tk.Label(root, text="Ready.", anchor="w", fg="#444")
        self.status_label.grid(row=6, column=0, columnspan=3, sticky="w", padx=12)

    def browse_input(self):
        folder = filedialog.askdirectory(title="Select Folder Containing .txt Files")
        if folder:
            self.input_var.set(folder)

    def browse_output(self):
        file = filedialog.asksaveasfilename(
            title="Save Excel File As",
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx"), ("All Files", "*.*")],
            initialfile="TxtFiles_Compiled.xlsx",
        )
        if file:
            self.output_var.set(file)

    def on_generate(self):
        input_folder = self.input_var.get().strip()
        output_file = self.output_var.get().strip()

        if not input_folder or not os.path.isdir(input_folder):
            messagebox.showerror("Invalid Input", "Please select a valid input folder.")
            return
        if not output_file:
            messagebox.showerror("Invalid Output", "Please choose an output Excel file.")
            return
        if not output_file.lower().endswith(".xlsx"):
            output_file += ".xlsx"
            self.output_var.set(output_file)

        self.generate_btn.config(state="disabled")
        self.progress["value"] = 0
        try:
            generate_excel(input_folder, output_file, self.status_label, self.progress)
        finally:
            self.generate_btn.config(state="normal")
            self.status_label.config(text="Ready.")
            self.progress["value"] = 0


if __name__ == "__main__":
    root = tk.Tk()
    app = TxtToExcelApp(root)
    root.mainloop()